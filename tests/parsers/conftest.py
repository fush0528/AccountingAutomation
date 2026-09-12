"""測試用的結算單建構工具。

測試不依賴 ``scripts/gen_fixtures.py`` 產生的檔案——那些是 gitignore 的，
CI 上不一定存在，而且內容會隨種子變動。測試要的是**固定不變的小樣本**，
所以在這裡就地組出位元組內容。

這就是 golden file 測試的精神：輸入固定、預期輸出固定，任何一方變動
都必須是有意識的決定。
"""

from __future__ import annotations

import io

import pytest

from reconciliation.parsers.base import SettlementSource

PLATFORM_A_HEADER = (
    "訂單編號,訂單成立日,商品名稱,數量,訂單金額,平台手續費,撥款金額,結算日期,交易類型"
)

PLATFORM_C_HEADER = "訂單序號,交易別,品名,數量,銷售金額,服務費,實撥金額,撥款日"

PLATFORM_B_HEADER = [
    "訂單編號",
    "商品",
    "數量",
    "訂單金額",
    "手續費",
    "撥款金額",
    "結算日",
    "類別",
]


def source(filename: str, text: str, encoding: str = "utf-8-sig") -> SettlementSource:
    return SettlementSource(filename=filename, content=text.encode(encoding))


@pytest.fixture
def platform_a_source() -> SettlementSource:
    rows = [
        PLATFORM_A_HEADER,
        "SP-24A7X9K2,2026-03-01,無線滑鼠,1,1000,55.30,944.70,2026-03-15,銷售",
        "sp-24b8y1m3,2026-03-02,機械鍵盤,1,2500,138.25,2361.75,2026-03-16,銷售",
        ",2026-03-03,USB 集線器,1,680,37.60,642.40,2026-03-17,銷售",
        "SP-24D1A3P5,2026-03-04,螢幕支架,2,2400,132.72,2267.28,2026-03-18,銷售",
        "SP-24A7X9K2,2026-03-01,無線滑鼠,1,-500,-27.65,-472.35,2026-03-20,退款",
    ]
    return source("platform_a_2026-03.csv", "\n".join(rows) + "\n")


@pytest.fixture
def platform_c_source() -> SettlementSource:
    rows = [
        PLATFORM_C_HEADER,
        'PC-24C9Z2N4,銷貨,USB集線器,1,"1,000","61.50","938.50",115/03/15',
        'PC-24E5B7Q9,銷貨,行動電源,2,"2,580","158.67","2,421.33",115/03/16',
        'PC-24F6C8R1,退貨,手機保護殼,1,"590","36.29","553.71",115/03/20',
        ',銷貨,充電線組,1,"460","28.29","431.71",115/03/21',
    ]
    return source("platform_c_2026-03.csv", "\r\n".join(rows) + "\r\n", encoding="big5")


def build_platform_b_workbook(data_rows: list[list[object]], *, title_rows: int = 3) -> bytes:
    """組出一份帶報表標題與兩層表頭的 Excel。

    ``title_rows`` 可調，用來驗證 parser 真的是「找」表頭而不是寫死列號。
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for i in range(title_rows):
        if i == 0:
            ws.cell(row=1, column=1, value="平台 B 商城　賣家結算明細表")
        elif i == 1:
            ws.cell(row=2, column=1, value="結算期間：2026/03/01 - 2026/03/31")
        # 其餘留白

    group_row = title_rows + 1
    ws.cell(row=group_row, column=1, value="訂單資訊")
    ws.merge_cells(start_row=group_row, start_column=1, end_row=group_row, end_column=3)
    ws.cell(row=group_row, column=4, value="金額明細")
    ws.merge_cells(start_row=group_row, start_column=4, end_row=group_row, end_column=6)

    header_row = group_row + 1
    for col, name in enumerate(PLATFORM_B_HEADER, start=1):
        ws.cell(row=header_row, column=col, value=name)

    for offset, values in enumerate(data_rows, start=header_row + 1):
        for col, value in enumerate(values, start=1):
            ws.cell(row=offset, column=col, value=value)

    ws.cell(row=header_row + len(data_rows) + 2, column=1, value="※ 本表僅供對帳參考")

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def platform_b_source() -> SettlementSource:
    rows: list[list[object]] = [
        ["MO-24A7X9K2", "無線滑鼠", 1, 1000, 48.0, 952.0, "2026/03/15", "銷售"],
        ["MO-24B8Y1M3", "機械鍵盤", 1, 2500.0, 120.0, 2380.0, "2026/03/16", "銷售"],
        ["", "USB 集線器", 1, 680, 32.64, 647.36, "2026/03/17", "銷售"],
        ["MO-24D1A3P5", "螢幕支架", 1, 1200, 57.6, 1142.4, "2026/03/18", "折讓"],
    ]
    return SettlementSource(
        filename="platform_b_2026-03.xlsx", content=build_platform_b_workbook(rows)
    )
