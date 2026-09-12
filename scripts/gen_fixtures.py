"""合成結算單產生器。

本專案沒有真實的平台結算單（作者手上沒有），所以測試資料由這支腳本產生。
這件事在 README 有明確說明，不做任何有真實資料的暗示。

但「合成」不等於「乾淨」。如果只產生完美對得起來的資料，對帳引擎就是在
解一個假問題——真實世界的困難全部來自髒資料。所以這支腳本會**刻意注入**
下列情況，每一種都對應對帳引擎必須處理的一個真實案例：

======================  ==================================================
注入的狀況              對帳引擎要怎麼處理
======================  ==================================================
訂單編號加前綴／大小寫  正規化後 Stage 1 仍可精確匹配
訂單編號整欄空白        Stage 1 找不到，落到 Stage 3 模糊匹配
金額不符（運費補貼）    Stage 2 容差匹配，計算差額並歸因
部分退款（獨立一列）    同一訂單出現多列，不能假設一對一
重複列（平台重複匯出）  匯入時的冪等性要擋掉
平台有、我方無          漏記單
我方有、平台無          未撥款（跨月結算或真的沒收到）
======================  ==================================================

同時輸出 ``manifest.json`` 作為 **ground truth**：裡面記錄了每一種狀況
實際注入了哪幾筆。W3 的 benchmark 可以拿它來驗證「自動匹配率 92%」這個
數字是真的算對了，而不是自己跟自己說好聽話。

執行::

    python scripts/gen_fixtures.py                 # 預設 300 筆訂單
    python scripts/gen_fixtures.py --orders 2000   # 壓力測試用
    python scripts/gen_fixtures.py --seed 42       # 換一組資料

產生的檔案在 ``data/generated/``，已被 .gitignore 排除——測試資料應該由
種子重建，不該進版控。
"""

from __future__ import annotations

import argparse
import csv
import json
import random
import sys
from dataclasses import asdict, dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

from reconciliation.domain import Money  # noqa: E402

OUTPUT_DIR = ROOT / "data" / "generated"

PRODUCTS = [
    ("無線滑鼠", "480", "1280"),
    ("機械鍵盤", "1580", "4200"),
    ("USB 集線器", "390", "980"),
    ("螢幕支架", "890", "2600"),
    ("筆電散熱座", "590", "1490"),
    ("藍牙耳機", "790", "3600"),
    ("行動電源", "450", "1290"),
    ("網路攝影機", "690", "2480"),
    ("手機保護殼", "180", "590"),
    ("充電線組", "120", "460"),
]

#: 各平台的手續費率。真實平台的費率會隨方案與品類變動，
#: 這裡固定住，讓「金額差異」只來自我們刻意注入的原因。
FEE_RATES = {
    "platform_a": Decimal("0.0553"),
    "platform_b": Decimal("0.0480"),
    "platform_c": Decimal("0.0615"),
}

PLATFORM_PREFIX = {
    "platform_a": "SP-",
    "platform_b": "MO-",
    "platform_c": "PC-",
}


# ----------------------------------------------------------------------
@dataclass
class Injected:
    """記錄實際注入了哪些狀況，寫進 manifest 作為 ground truth。"""

    missing_order_id: list[str] = field(default_factory=list)
    amount_variance: list[str] = field(default_factory=list)
    partial_refund: list[str] = field(default_factory=list)
    duplicated_row: list[str] = field(default_factory=list)
    settlement_only: list[str] = field(default_factory=list)
    ledger_only: list[str] = field(default_factory=list)
    lowercase_id: list[str] = field(default_factory=list)


@dataclass
class Order:
    order_id: str
    platform: str
    external_order_id: str
    ordered_at: date
    product_name: str
    quantity: int
    gross: Money
    fee: Money

    @property
    def net(self) -> Money:
        return self.gross - self.fee


@dataclass
class Row:
    """準備寫進結算單的一列（還沒套用各平台的格式）。"""

    external_order_id: str
    product_name: str
    quantity: int
    gross: Money
    fee: Money
    net: Money
    settled_at: date
    kind: str  # sale / refund / adjustment


# ----------------------------------------------------------------------
def make_orders(rng: random.Random, count: int) -> list[Order]:
    orders: list[Order] = []
    start = date(2026, 3, 1)
    platforms = list(FEE_RATES)
    for i in range(1, count + 1):
        platform = rng.choice(platforms)
        name, low, high = rng.choice(PRODUCTS)
        unit = Money.from_str(
            str(
                rng.randint(
                    int(Money.from_str(low).minor_units / 100),
                    int(Money.from_str(high).minor_units / 100),
                )
            )
        )
        quantity = rng.choices([1, 1, 1, 2, 2, 3, 5], k=1)[0]
        gross = unit * quantity
        fee = gross * FEE_RATES[platform]
        code = f"{rng.randint(0x10000000, 0xFFFFFFFF):08X}"
        orders.append(
            Order(
                order_id=f"ORD-{i:05d}",
                platform=platform,
                external_order_id=code,
                ordered_at=start + timedelta(days=rng.randint(0, 27)),
                product_name=name,
                quantity=quantity,
                gross=gross,
                fee=fee,
            )
        )
    return orders


def build_rows(rng: random.Random, orders: list[Order], injected: Injected) -> dict[str, list[Row]]:
    """把訂單轉成各平台的結算列，並沿途注入髒資料。"""
    rows: dict[str, list[Row]] = {p: [] for p in FEE_RATES}

    for order in orders:
        settled = order.ordered_at + timedelta(days=rng.randint(10, 20))
        roll = rng.random()

        # 5%：我方有、平台無（跨月結算或真的沒撥款）
        if roll < 0.05:
            injected.ledger_only.append(order.order_id)
            continue

        gross, fee, net = order.gross, order.fee, order.net
        external = PLATFORM_PREFIX[order.platform] + order.external_order_id

        # 6%：訂單編號整欄空白 → 只能靠 Stage 3 救回來
        if 0.05 <= roll < 0.11:
            injected.missing_order_id.append(order.order_id)
            external = ""
        # 5%：編號被寫成小寫 → 正規化後 Stage 1 仍應命中
        elif 0.11 <= roll < 0.16:
            injected.lowercase_id.append(order.order_id)
            external = external.lower()

        # 7%：金額與我方預期不符。
        #
        # 注意注入的方式：結算列**本身仍然自洽**（gross - fee == net），
        # 差異只存在於「平台說的」與「我方預期的」之間。真實的結算單不會
        # 自己算錯加減法，差異來自費率調整或運費補貼這類業務原因。
        # 如果直接去改 net 讓該列算術不成立，那是在模擬「parser 壞掉」，
        # 不是在模擬「需要對帳的差異」——兩者是完全不同的問題。
        if rng.random() < 0.07:
            if rng.random() < 0.5:
                # 費率調整：平台實收的手續費與我方預期不同
                fee = fee + Money.from_str(str(rng.choice([-40, -25, 15, 30])))
            else:
                # 運費補貼或加購：平台認列的訂單金額與我方不同
                gross = gross + Money.from_str(str(rng.choice([-60, 45, 80])))
            net = gross - fee
            injected.amount_variance.append(order.order_id)

        rows[order.platform].append(
            Row(external, order.product_name, order.quantity, gross, fee, net, settled, "sale")
        )

        # 4%：部分退款，獨立成一列 → 同一訂單出現兩列
        if rng.random() < 0.04:
            ratio = Decimal(str(rng.choice([0.3, 0.5, 1.0])))
            r_gross = -(gross * ratio)
            r_fee = -(fee * ratio)
            injected.partial_refund.append(order.order_id)
            rows[order.platform].append(
                Row(
                    external,
                    order.product_name,
                    order.quantity,
                    r_gross,
                    r_fee,
                    r_gross - r_fee,
                    settled + timedelta(days=rng.randint(1, 8)),
                    "refund",
                )
            )

    # 3%：平台有、我方無（漏記單，或平台自行產生的調整列）
    for platform, platform_rows in rows.items():
        for _ in range(max(1, int(len(platform_rows) * 0.03))):
            name, low, _high = rng.choice(PRODUCTS)
            gross = Money.from_str(low)
            fee = gross * FEE_RATES[platform]
            code = PLATFORM_PREFIX[platform] + f"{rng.randint(0x10000000, 0xFFFFFFFF):08X}"
            injected.settlement_only.append(code)
            platform_rows.append(
                Row(
                    code,
                    name,
                    1,
                    gross,
                    fee,
                    gross - fee,
                    date(2026, 3, rng.randint(10, 28)),
                    "sale",
                )
            )

    # 2%：平台重複匯出同一列 → 匯入的冪等性要擋掉
    for platform_rows in rows.values():
        candidates = [r for r in platform_rows if r.external_order_id]
        for row in rng.sample(candidates, k=max(1, int(len(candidates) * 0.02))):
            injected.duplicated_row.append(row.external_order_id)
            platform_rows.append(row)

    for platform_rows in rows.values():
        rng.shuffle(platform_rows)
    return rows


# ----------------------------------------------------------------------
def write_orders_csv(path: Path, orders: list[Order]) -> None:
    """我方的訂單資料。對帳時的另一半。"""
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "訂單編號",
                "平台",
                "平台訂單編號",
                "訂單時間",
                "商品名稱",
                "數量",
                "訂單金額",
                "預期手續費",
            ]
        )
        for o in orders:
            writer.writerow(
                [
                    o.order_id,
                    o.platform,
                    PLATFORM_PREFIX[o.platform] + o.external_order_id,
                    o.ordered_at.isoformat(),
                    o.product_name,
                    o.quantity,
                    o.gross.to_decimal(),
                    o.fee.to_decimal(),
                ]
            )


def write_platform_a(path: Path, rows: list[Row]) -> None:
    """UTF-8 CSV，單層表頭，西元日期。"""
    labels = {"sale": "銷售", "refund": "退款", "adjustment": "調整"}
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(
            [
                "訂單編號",
                "訂單成立日",
                "商品名稱",
                "數量",
                "訂單金額",
                "平台手續費",
                "撥款金額",
                "結算日期",
                "交易類型",
            ]
        )
        for r in rows:
            writer.writerow(
                [
                    r.external_order_id,
                    (r.settled_at - timedelta(days=14)).isoformat(),
                    r.product_name,
                    r.quantity,
                    r.gross.to_decimal(),
                    r.fee.to_decimal(),
                    r.net.to_decimal(),
                    r.settled_at.isoformat(),
                    labels[r.kind],
                ]
            )


def write_platform_b(path: Path, rows: list[Row]) -> None:
    """Excel，報表標題 + 兩層表頭 + 合併儲存格 + 數值型金額。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    labels = {"sale": "銷售", "refund": "折讓", "adjustment": "調整"}
    wb = Workbook()
    ws = wb.active
    ws.title = "結算明細"

    ws["A1"] = "平台 B 商城　賣家結算明細表"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")
    ws["A2"] = "結算期間：2026/03/01 - 2026/03/31"
    ws.merge_cells("A2:H2")
    # 第 3 列刻意留白——真實報表就長這樣

    ws["A4"] = "訂單資訊"
    ws.merge_cells("A4:C4")
    ws["D4"] = "金額明細"
    ws.merge_cells("D4:F4")
    ws["G4"] = "結算"
    ws.merge_cells("G4:H4")

    header = ["訂單編號", "商品", "數量", "訂單金額", "手續費", "撥款金額", "結算日", "類別"]
    for col, name in enumerate(header, start=1):
        cell = ws.cell(row=5, column=col, value=name)
        cell.font = Font(bold=True)

    for offset, r in enumerate(rows, start=6):
        ws.cell(row=offset, column=1, value=r.external_order_id)
        ws.cell(row=offset, column=2, value=r.product_name)
        ws.cell(row=offset, column=3, value=r.quantity)
        # 金額寫成數值而非字串——parser 必須處理 float/int
        ws.cell(row=offset, column=4, value=float(r.gross.to_decimal()))
        ws.cell(row=offset, column=5, value=float(r.fee.to_decimal()))
        ws.cell(row=offset, column=6, value=float(r.net.to_decimal()))
        ws.cell(row=offset, column=7, value=r.settled_at.strftime("%Y/%m/%d"))
        ws.cell(row=offset, column=8, value=labels[r.kind])

    ws.cell(row=len(rows) + 7, column=1, value="※ 本表僅供對帳參考")
    wb.save(path)


def write_platform_c(path: Path, rows: list[Row]) -> None:
    """Big5 CSV，民國年，千分位金額，退貨列印成正數。"""
    labels = {"sale": "銷貨", "refund": "退貨", "adjustment": "調整"}

    def roc(d: date) -> str:
        return f"{d.year - 1911}/{d.month:02d}/{d.day:02d}"

    def amount(m: Money) -> str:
        return f"{abs(m.to_decimal()):,}"

    lines = [["訂單序號", "交易別", "品名", "數量", "銷售金額", "服務費", "實撥金額", "撥款日"]]
    for r in rows:
        lines.append(
            [
                r.external_order_id,
                labels[r.kind],
                r.product_name,
                str(r.quantity),
                amount(r.gross),
                amount(r.fee),
                amount(r.net),
                roc(r.settled_at),
            ]
        )

    buffer = []
    for line in lines:
        buffer.append(",".join(f'"{c}"' if "," in c else c for c in line))
    text = "\r\n".join(buffer) + "\r\n"
    # errors="replace" 是刻意的：Big5 表達不了某些字元，真實檔案就會這樣
    path.write_bytes(text.encode("big5", errors="replace"))


# ----------------------------------------------------------------------
def main() -> int:
    parser = argparse.ArgumentParser(description="產生合成的結算單測試資料")
    parser.add_argument("--orders", type=int, default=300, help="訂單筆數（預設 300）")
    parser.add_argument("--seed", type=int, default=20260908, help="亂數種子，決定資料內容")
    parser.add_argument("--out", type=Path, default=OUTPUT_DIR, help="輸出目錄")
    args = parser.parse_args()

    rng = random.Random(args.seed)
    args.out.mkdir(parents=True, exist_ok=True)

    orders = make_orders(rng, args.orders)
    injected = Injected()
    rows = build_rows(rng, orders, injected)

    write_orders_csv(args.out / "orders.csv", orders)
    write_platform_a(args.out / "platform_a_2026-03.csv", rows["platform_a"])
    write_platform_b(args.out / "platform_b_2026-03.xlsx", rows["platform_b"])
    write_platform_c(args.out / "platform_c_2026-03.csv", rows["platform_c"])

    manifest = {
        "seed": args.seed,
        "order_count": len(orders),
        "settlement_rows": {p: len(r) for p, r in rows.items()},
        "injected": {k: len(v) for k, v in asdict(injected).items()},
        "injected_detail": asdict(injected),
        "note": "合成資料。injected_detail 是 ground truth，供 W3 驗證匹配率用。",
    }
    (args.out / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    print(f"輸出目錄：{args.out}")
    print(f"  亂數種子 {args.seed}（同一個種子永遠產生同一份資料）\n")
    print(f"  orders.csv                    我方訂單 {len(orders)} 筆")
    for platform, platform_rows in sorted(rows.items()):
        print(f"  {platform}_2026-03.*{' ' * 9}結算列 {len(platform_rows)} 筆")
    print("\n  刻意注入的狀況：")
    for name, items in sorted(asdict(injected).items()):
        print(f"    {name:<20} {len(items):>4} 筆")
    print("\n  manifest.json 記錄了上表的明細，作為 W3 驗證匹配率的 ground truth。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
